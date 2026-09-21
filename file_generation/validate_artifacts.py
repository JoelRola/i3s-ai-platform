from pathlib import Path
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from fpdf import FPDF
from PIL import Image
import csv
import matplotlib.pyplot as plt

out = Path.home() / 'file-generation-validation'
out.mkdir(parents=True, exist_ok=True)

doc = Document(); doc.add_heading('Test Report', 0)
doc.add_paragraph('First validation paragraph.'); doc.add_paragraph('Second validation paragraph.')
t = doc.add_table(rows=3, cols=3)
for r in range(3):
    for c in range(3): t.cell(r,c).text = f'R{r+1}C{c+1}'
doc.save(out/'test_report.docx')

wb = Workbook(); ws = wb.active; ws.title = 'Finance'
ws.append(['Month','Revenue','Expenses','Profit'])
for i, month in enumerate(['Jan','Feb','Mar','Apr','May'], 2):
    ws.append([month, i*1000, i*400, f'=B{i}-C{i}'])
wb.save(out/'test_finance.xlsx')

prs = Presentation()
for title in ['Overview','Results','Next Steps']:
    slide = prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text = title
prs.save(out/'test_presentation.pptx')

pdf = FPDF(); pdf.add_page(); pdf.set_font('Helvetica', size=14); pdf.cell(text='Test Summary')
pdf.output(out/'test_summary.pdf')

with (out/'test_data.csv').open('w', newline='') as f:
    w = csv.writer(f); w.writerow(['Month','Revenue']); w.writerows([['Jan',1000],['Feb',1200],['Mar',1400]])

plt.plot(['Jan','Feb','Mar'], [1000,1200,1400], marker='o'); plt.title('Revenue'); plt.tight_layout(); plt.savefig(out/'revenue_chart.png'); plt.close()

# Validate formats and formulas after reopening.
assert (out/'test_report.docx').exists()
assert load_workbook(out/'test_finance.xlsx', data_only=False).active['D2'].value == '=B2-C2'
assert len(Presentation(out/'test_presentation.pptx').slides) == 3
assert (out/'test_summary.pdf').read_bytes().startswith(b'%PDF-')
assert Image.open(out/'revenue_chart.png').format == 'PNG'
for p in sorted(out.iterdir()): print(f'{p.name}|{p.stat().st_size}')
